import urllib.request
import urllib.parse
import json
import io
import openpyxl

BASE_URL = "http://127.0.0.1:8000"

def test_html_and_js():
    print("Testing HTML and JS content...")
    # Test HTML
    req = urllib.request.urlopen(f"{BASE_URL}/")
    html = req.read().decode("utf-8")
    
    assert "corrTableFilterOp" in html, "Missing corrTableFilterOp in HTML"
    assert "corrTableFilterVal" in html, "Missing corrTableFilterVal in HTML"
    assert "resetCorrTableFilterBtn" in html, "Missing resetCorrTableFilterBtn in HTML"
    assert "Экспорт в Excel" not in html, "'Экспорт в Excel' should be removed, replaced with 'Экспорт'"
    
    # Check button texts
    assert 'id="exportMatrixCorrExcelBtn"' in html
    assert 'id="exportMatrixBetaExcelBtn"' in html
    assert 'id="corrExportExcelBtn"' in html
    print("[OK] HTML verification passed")

    # Test JS
    req_js = urllib.request.urlopen(f"{BASE_URL}/static/app.js")
    js = req_js.read().decode("utf-8")
    assert "function parseFilterNumber" in js, "Missing parseFilterNumber in app.js"
    assert "matchingSet" in js, "Missing mutual matchingSet in app.js"
    assert "diagValid" in js, "Missing diagValid in app.js"
    assert "lastFilteredCorrList" in js, "Missing lastFilteredCorrList in app.js"
    print("[OK] JS verification passed")

def test_corr_filtered_export():
    print("Testing filtered correlation table export...")
    payload = {
        "period": "1y",
        "target_move": 2.5,
        "rows": [
            {
                "symbol": "SPY",
                "type": "Benchmark",
                "correlation": 0.85,
                "beta": 1.10,
                "volatility": 0.18,
                "observations": 252,
                "is_partial": False,
                "_peerMove": 2.0,
                "_expMove": 2.2,
                "_diff": 0.3,
                "_normDiff": 1.67
            },
            {
                "symbol": "QQQ",
                "type": "Benchmark",
                "correlation": 0.78,
                "beta": 1.25,
                "volatility": 0.22,
                "observations": 252,
                "is_partial": False,
                "_peerMove": 2.0,
                "_expMove": 2.5,
                "_diff": 0.0,
                "_normDiff": 0.0
            }
        ]
    }
    
    data_bytes = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/api/ticker/AAPL/correlations/export",
        data=data_bytes,
        headers={"Content-Type": "application/json"}
    )
    res = urllib.request.urlopen(req)
    assert res.status == 200, f"Export returned status {res.status}"
    
    content = res.read()
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    
    # Rows should only contain the 2 filtered assets: SPY and QQQ
    tickers_in_excel = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0]:
            tickers_in_excel.append(row[0])
    
    assert tickers_in_excel == ["SPY", "QQQ"], f"Expected ['SPY', 'QQQ'], got {tickers_in_excel}"
    print(f"[OK] Filtered correlation export verified: only exported rows {tickers_in_excel}")

def test_matrix_beta_export_with_masking():
    print("Testing Beta Matrix export with cell masking (Beta threshold >= 5)...")
    symbols = ["BRUN", "HQ", "XNDU"]
    # Suppose BRUN and HQ have beta 0.296 (not >= 5, so masked as null)
    # BRUN and XNDU have beta 5.2 (>= 5, so included)
    matrix = {
        "BRUN": {"BRUN": None, "HQ": None, "XNDU": 5.210},
        "HQ": {"BRUN": None, "HQ": None, "XNDU": None},
        "XNDU": {"BRUN": 5.150, "HQ": None, "XNDU": None}
    }
    payload = {
        "type": "beta",
        "title": "Matrix Beta (formul Slope)",
        "period": "1y",
        "symbols": symbols,
        "matrix": matrix
    }
    
    data_bytes = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/api/matrix/export",
        data=data_bytes,
        headers={"Content-Type": "application/json"}
    )
    res = urllib.request.urlopen(req)
    assert res.status == 200, f"Beta Matrix export returned {res.status}"
    
    wb = openpyxl.load_workbook(io.BytesIO(res.read()))
    ws = wb.active
    
    # Check row 3 (BRUN): Col 1 is BRUN, Col 2 (BRUN) is —, Col 3 (HQ) is —, Col 4 (XNDU) is 5.21
    brun_row = [ws.cell(row=3, column=c).value for c in range(1, 5)]
    assert brun_row[0] == "BRUN"
    assert brun_row[1] == "—", f"Diagonal should be masked, got {brun_row[1]}"
    assert brun_row[2] == "—", f"Non-matching cell should be masked, got {brun_row[2]}"
    assert brun_row[3] == 5.21, f"Matching cell should be 5.21, got {brun_row[3]}"
    print("[OK] Beta matrix export properly masks non-matching cells and diagonal as '—'")

def test_matrix_corr_export_with_masking():
    print("Testing Pearson Matrix export with masking (r >= 0.7)...")
    symbols = ["BRUN", "CORZ", "CRWV"]
    matrix = {
        "BRUN": {"BRUN": 1.0, "CORZ": None, "CRWV": 0.725},
        "CORZ": {"BRUN": None, "CORZ": 1.0, "CRWV": 0.717},
        "CRWV": {"BRUN": 0.725, "CORZ": 0.717, "CRWV": 1.0}
    }
    payload = {
        "type": "correlation",
        "title": "Matrix Correlation Pearson",
        "period": "1y",
        "symbols": symbols,
        "matrix": matrix
    }
    
    data_bytes = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/api/matrix/export",
        data=data_bytes,
        headers={"Content-Type": "application/json"}
    )
    res = urllib.request.urlopen(req)
    assert res.status == 200, f"Corr Matrix export returned {res.status}"
    
    wb = openpyxl.load_workbook(io.BytesIO(res.read()))
    ws = wb.active
    
    # Check row 3 (BRUN): Col 1 is BRUN, Col 2 (BRUN) is 1.0, Col 3 (CORZ) is —, Col 4 (CRWV) is 0.725
    brun_row = [ws.cell(row=3, column=c).value for c in range(1, 5)]
    assert brun_row[0] == "BRUN"
    assert brun_row[1] == 1.0, f"Diagonal should be 1.0, got {brun_row[1]}"
    assert brun_row[2] == "—", f"Non-matching cell should be '—', got {brun_row[2]}"
    assert brun_row[3] == 0.725, f"Matching cell should be 0.725, got {brun_row[3]}"
    print("[OK] Pearson matrix export properly masks non-matching cells as '—' and shows matching")

if __name__ == "__main__":
    test_html_and_js()
    test_corr_filtered_export()
    test_matrix_beta_export_with_masking()
    test_matrix_corr_export_with_masking()
    print("\nALL TESTS PASSED SUCCESSFULLY!")
