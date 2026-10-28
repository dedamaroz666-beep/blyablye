import urllib.request
import json
import io
import openpyxl

BASE_URL = "http://127.0.0.1:8000"

def test_user_request():
    with urllib.request.urlopen(f"{BASE_URL}/", timeout=10) as resp:
        html = resp.read().decode("utf-8")

    # 1. No emojis in Tier view
    tier_chunk = html[html.find('id="viewTier"'):html.find('id="datalensFilterModalBackdrop"')]
    for emoji in ["⚡", "📄", "📊", "📐", "⚖️", "⚖"]:
        assert emoji not in tier_chunk, f"Emoji {emoji} still present in viewTier HTML!"
    print("[PASS] 1. Emojis removed from Tier view buttons and empty state.")

    # 2. Main heading and subheading removed
    assert "Tier Position Sizing & Risk Parity" not in html, "Heading 'Tier Position Sizing & Risk Parity' still present!"
    assert "Расчёт лота и риск-паритета позиций" not in html, "Subheading still present!"
    print("[PASS] 2. Heading and subheading removed.")

    # 3. Screenshot 2 note removed
    assert "Формула оценивает Open, High, Low, Close каждого дня." not in html, "Screenshot 2 note still present!"
    print("[PASS] 3. Formula explanation note from screenshot 2 removed.")

    # 4. Table headers check
    assert "<th>ЦЕНА ($)</th>" not in html
    assert "<th>GK ВОЛ (ДЕНЬ)</th>" not in html
    assert "<th>GK ВОЛ (ГОД)</th>" not in html
    assert "<th>РИСК 1 АКЦ. ($)</th>" not in html
    assert "<th>ТИР 2 (2X)</th>" not in html
    assert "<th>ТИР 3 (3X)</th>" not in html
    assert "ТИКЕР</th>" in html
    assert "ТИР 1 (АКЦИИ)" in html
    assert "КАПИТАЛ 1 ТИР</th>" in html
    assert "РИСК 1 ТИР</th>" in html
    assert "ПАРИТЕТ</th>" in html
    print("[PASS] 4. Table columns in UI contain only Ticker, Tier 1 Shares, Capital, Risk, Parity.")

    # 5. Excel Export verification
    payload = {
        "benchmark": "MU",
        "benchmark_shares": 5,
        "lookback_days": 20,
        "rounding": "1",
        "tickers": ["MRVL", "NVDA"]
    }
    req = urllib.request.Request(
        f"{BASE_URL}/api/tier/export/excel",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        wb = openpyxl.load_workbook(io.BytesIO(resp.read()))
        ws = wb.active
        assert ws.max_column == 2, f"Excel should have only 2 columns, got {ws.max_column}"
        h1 = ws.cell(row=1, column=1).value
        h2 = ws.cell(row=1, column=2).value
        assert h1 == "ТИКЕР", f"Col 1 should be ТИКЕР, got {h1}"
        assert h2 == "ТИР 1 (АКЦИИ)", f"Col 2 should be ТИР 1 (АКЦИИ), got {h2}"
        row2 = [ws.cell(row=2, column=c).value for c in range(1, 3)]
        assert row2[0] == "MRVL", f"Row 2 ticker mismatch: {row2}"
        assert isinstance(row2[1], (int, float)) and row2[1] > 0, f"Row 2 shares invalid: {row2}"
        print(f"[PASS] 5. Excel export has exactly 2 columns: {h1} and {h2}, first row: {row2}")

    print("\nALL USER REQUIREMENTS VERIFIED SUCCESSFULLY!")

if __name__ == "__main__":
    test_user_request()
