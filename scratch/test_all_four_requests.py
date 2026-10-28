import urllib.request
import urllib.error
import json
import io
import openpyxl

def test_all():
    print("--- 1. Testing Pearson Matrix Excel Export ---")
    corr_payload = {
        "type": "correlation",
        "title": "Matrix Correlation Pearson",
        "period": "1y",
        "symbols": ["AAPL", "MSFT", "NVDA", "SQQQ"],
        "matrix": {
            "AAPL": {"AAPL": 1.0, "MSFT": 0.72, "NVDA": 0.45, "SQQQ": -0.65},
            "MSFT": {"AAPL": 0.72, "MSFT": 1.0, "NVDA": 0.55, "SQQQ": -0.58},
            "NVDA": {"AAPL": 0.45, "MSFT": 0.55, "NVDA": 1.0, "SQQQ": -0.71},
            "SQQQ": {"AAPL": -0.65, "MSFT": -0.58, "NVDA": -0.71, "SQQQ": 1.0}
        }
    }
    req = urllib.request.Request(
        "http://127.0.0.1:8000/api/matrix/export",
        data=json.dumps(corr_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.read()))
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "Тикер"
        assert [ws.cell(row=2, column=c).value for c in range(2, 6)] == ["AAPL", "MSFT", "NVDA", "SQQQ"]
        # AAPL-AAPL diagonal (1.000)
        assert ws.cell(row=3, column=2).value == 1.0
        # High corr fill (AAPL-MSFT = 0.72)
        c_fill = ws.cell(row=3, column=3).fill.start_color.rgb
        assert c_fill in ["00DCFCE7", "DCFCE7"], f"Unexpected high corr fill: {c_fill}"
        # Neg corr fill (AAPL-SQQQ = -0.65)
        neg_fill = ws.cell(row=3, column=5).fill.start_color.rgb
        assert neg_fill in ["00FEE2E2", "FEE2E2"], f"Unexpected neg corr fill: {neg_fill}"
    print("Pearson Matrix Excel export: PASSED")

    print("\n--- 2. Testing Beta Matrix Excel Export (Photo 1 colors) ---")
    beta_payload = {
        "type": "beta",
        "title": "Matrix Beta (formul Slope)",
        "period": "1y",
        "symbols": ["SPY", "QQQ", "AAPL", "SQQQ"],
        "matrix": {
            "SPY": {"SPY": 1.0, "QQQ": 0.95, "AAPL": 0.65, "SQQQ": -2.50},
            "QQQ": {"SPY": 1.05, "QQQ": 1.0, "AAPL": 1.35, "SQQQ": -2.80},
            "AAPL": {"SPY": 1.25, "QQQ": 1.10, "AAPL": 1.0, "SQQQ": -1.85},
            "SQQQ": {"SPY": -2.20, "QQQ": -2.60, "AAPL": -1.50, "SQQQ": 1.0}
        }
    }
    req = urllib.request.Request(
        "http://127.0.0.1:8000/api/matrix/export",
        data=json.dumps(beta_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.read()))
        ws = wb.active
        assert ws.cell(row=2, column=1).value == "Тикер"
        # Diagonal SPY-SPY: soft blue (00BFDBFE)
        diag_fill = ws.cell(row=3, column=2).fill.start_color.rgb
        assert diag_fill in ["00BFDBFE", "BFDBFE"], f"Unexpected diag fill: {diag_fill}"
        # SPY-QQQ = 0.95 (market beta 0.85-1.15: soft sky blue 00BFDBFE)
        b_neutral_fill = ws.cell(row=3, column=3).fill.start_color.rgb
        assert b_neutral_fill in ["00BFDBFE", "BFDBFE"], f"Unexpected beta neutral fill: {b_neutral_fill}"
        # SPY-AAPL = 0.65 (low beta < 0.85: soft warm cream 00FEF3C7)
        b_low_fill = ws.cell(row=3, column=4).fill.start_color.rgb
        assert b_low_fill in ["00FEF3C7", "FEF3C7"], f"Unexpected beta low fill: {b_low_fill}"
        # SPY-SQQQ = -2.50 (negative beta < 0: soft pink/salmon 00FECACA)
        b_neg_fill = ws.cell(row=3, column=5).fill.start_color.rgb
        assert b_neg_fill in ["00FECACA", "FECACA"], f"Unexpected beta neg fill: {b_neg_fill}"
        # QQQ-AAPL = 1.35 (high beta > 1.15: soft mint green 00BBF7D0)
        b_high_fill = ws.cell(row=4, column=4).fill.start_color.rgb
        assert b_high_fill in ["00BBF7D0", "BBF7D0"], f"Unexpected beta high fill: {b_high_fill}"
    print("Beta Matrix Excel export (Photo 1 palette): PASSED")

    print("\n--- 3. Testing Ticker Correlations Table Export (Header & Fills) ---")
    corr_tbl_payload = {
        "period": "1y",
        "target_move": 2.0,
        "rows": [
            {
                "symbol": "SPY",
                "type": "Benchmark",
                "correlation": 0.85,
                "beta": 1.05,
                "observations": 252
            }
        ]
    }
    req = urllib.request.Request(
        "http://127.0.0.1:8000/api/ticker/AAPL/correlations/export",
        data=json.dumps(corr_tbl_payload).encode("utf-8"),
        headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200
        wb = openpyxl.load_workbook(io.BytesIO(resp.read()))
        ws = wb.active
        # Header 1 is "Тикер" instead of "Актив / Бенчмарк"
        assert ws.cell(row=2, column=1).value == "Тикер", f"Col 1 header was: {ws.cell(row=2, column=1).value}"
    print("Ticker Correlations Table Export (with 'Тикер'): PASSED")

    print("\n--- 4. Testing HTML Subtitles Removal & Buttons ---")
    with urllib.request.urlopen("http://127.0.0.1:8000/") as resp:
        html = resp.read().decode("utf-8")
        # Ensure earnings subtitle is removed
        assert "Динамика цен акций в окнах R-1" not in html, "Earnings subtitle still in HTML!"
        # Ensure dividends subtitle is removed
        assert "Даты отсечки, размер дивиденда" not in html, "Dividends subtitle still in HTML!"
        # Ensure matrix export buttons exist
        assert 'id="exportMatrixCorrExcelBtn"' in html, "exportMatrixCorrExcelBtn missing in HTML"
        assert 'id="exportMatrixBetaExcelBtn"' in html, "exportMatrixBetaExcelBtn missing in HTML"
        # Ensure corrTable has "Тикер"
        assert '<th data-sort="symbol" class="th-ticker-col">Тикер' in html, "corrTable th-ticker-col missing"
    print("HTML Subtitles Removal & Buttons: PASSED")

    print("\n--- 5. Testing CSS Column Sizing ---")
    with urllib.request.urlopen("http://127.0.0.1:8000/static/style.css") as resp:
        css = resp.read().decode("utf-8")
        assert ".th-ticker-col" in css and ".td-ticker-col" in css, "Ticker col CSS missing in style.css"
        assert "width: 95px" in css, "Narrow width missing in style.css"
    print("CSS Column Sizing: PASSED")

    print("\n--- 6. Testing JS Color Palettes & Export Handlers ---")
    with urllib.request.urlopen("http://127.0.0.1:8000/static/app.js") as resp:
        js = resp.read().decode("utf-8")
        assert "rgba(254, 202, 202" in js, "Pink beta color missing in app.js"
        assert "rgba(254, 243, 199" in js, "Cream beta color missing in app.js"
        assert "rgba(191, 219, 254" in js, "Sky blue beta color missing in app.js"
        assert "rgba(187, 247, 208" in js, "Mint green beta color missing in app.js"
        assert "exportMatrixCorrelationToExcel" in js, "exportMatrixCorrelationToExcel missing in app.js"
        assert "exportMatrixBetaToExcel" in js, "exportMatrixBetaToExcel missing in app.js"
        assert "td-ticker-col" in js, "td-ticker-col missing in app.js"
    print("JS Color Palettes & Export Handlers: PASSED")

    print("\n=== ALL 4 USER REQUESTS VERIFIED AND PASSED SUCCESSFULLY! ===")

if __name__ == "__main__":
    test_all()
