import urllib.request
import urllib.error
import json
import io
import openpyxl

def test_excel_export():
    print("Testing Excel export endpoint...")
    url = "http://127.0.0.1:8000/api/ticker/AAPL/correlations/export"
    payload = {
        "period": "1y",
        "target_move": 2.5,
        "rows": [
            {
                "symbol": "SPY",
                "type": "Benchmark",
                "correlation": 0.812,
                "beta": 1.15,
                "_peerMove": 1.0,
                "_expMove": 1.15,
                "_targetMove": 2.5,
                "_diff": 1.35,
                "_normDiff": 0.85,
                "observations": 252,
                "is_partial": False
            },
            {
                "symbol": "QQQ",
                "type": "Benchmark",
                "correlation": 0.745,
                "beta": 0.95,
                "_peerMove": 1.0,
                "_expMove": 0.95,
                "_targetMove": 2.5,
                "_diff": 1.55,
                "_normDiff": 0.98,
                "observations": 252,
                "is_partial": False
            },
            {
                "symbol": "RDDT",
                "type": "Peer",
                "correlation": -0.325,
                "beta": -0.929,
                "_peerMove": 1.0,
                "_expMove": -0.929,
                "_targetMove": 2.5,
                "_diff": 3.429,
                "_normDiff": 1.45,
                "observations": 120,
                "is_partial": True,
                "partial_reason": "Доступно 120 дн. из 252 дн."
            },
            {
                "symbol": "SQQQ",
                "type": "Peer",
                "correlation": -0.755,
                "beta": -2.85,
                "_peerMove": 1.0,
                "_expMove": -2.85,
                "_targetMove": 2.5,
                "_diff": 5.35,
                "_normDiff": 2.10,
                "observations": 252,
                "is_partial": False
            }
        ]
    }
    
    data_bytes = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=data_bytes,
        headers={"Content-Type": "application/json"}
    )
    
    with urllib.request.urlopen(req) as resp:
        assert resp.status == 200, f"Expected 200, got {resp.status}"
        ctype = resp.headers.get("Content-Type")
        print(f"Content-Type: {ctype}")
        assert "spreadsheetml.sheet" in ctype, f"Invalid content type: {ctype}"
        content_disp = resp.headers.get("Content-Disposition")
        print(f"Content-Disposition: {content_disp}")
        assert "AAPL_correlation_matrix_1y.xlsx" in content_disp, f"Unexpected header: {content_disp}"
        
        file_bytes = resp.read()
        print(f"Excel file size: {len(file_bytes)} bytes")
        assert len(file_bytes) > 2000, "Excel file is suspiciously small"
        
        # Load workbook to verify contents and styling
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
        sheet = wb.active
        print(f"Active sheet: {sheet.title}")
        assert "AAPL Correlations" in sheet.title
        
        # Check title
        assert "CHEPUH / FinWhale" in sheet["A1"].value
        
        # Check headers at row 2
        headers = [sheet.cell(row=2, column=col).value for col in range(1, 12)]
        print(f"Headers: {headers}")
        assert "Актив / Бенчмарк" in headers
        assert "Корреляция" in headers
        assert "Beta" in headers
        assert "Разница (Diff %)" in headers
        assert "Статус истории" in headers
        
        # Check rows
        symbols = [sheet.cell(row=r, column=1).value for r in range(3, 7)]
        print(f"Exported symbols in sheet: {symbols}")
        assert symbols == ["SPY", "QQQ", "RDDT", "SQQQ"]
        
        # Check SPY positive correlation and beta fill
        spy_corr_fill = sheet.cell(row=3, column=3).fill.start_color.rgb
        spy_beta_fill = sheet.cell(row=3, column=4).fill.start_color.rgb
        print(f"SPY corr fill: {spy_corr_fill}, beta fill: {spy_beta_fill}")
        assert spy_corr_fill == "00E8F5E9" or spy_corr_fill == "E8F5E9", f"Unexpected fill {spy_corr_fill}"
        
        # Check RDDT negative correlation and beta fill
        rddt_corr_fill = sheet.cell(row=5, column=3).fill.start_color.rgb
        rddt_beta_fill = sheet.cell(row=5, column=4).fill.start_color.rgb
        print(f"RDDT corr fill: {rddt_corr_fill}, beta fill: {rddt_beta_fill}")
        assert rddt_corr_fill == "00FFEBEE" or rddt_corr_fill == "FFEBEE", f"Unexpected fill {rddt_corr_fill}"
        
        # Check partial warning for RDDT
        rddt_part_status = sheet.cell(row=5, column=11).value
        print(f"RDDT status: {rddt_part_status}")
        assert "120 дн." in rddt_part_status

    print("Excel export verification PASSED!\n")


def test_html_and_js_served():
    print("Testing HTML and static assets served by server...")
    # HTML
    with urllib.request.urlopen("http://127.0.0.1:8000/") as resp:
        html = resp.read().decode("utf-8")
        assert 'id="corrExportExcelBtn"' in html, "corrExportExcelBtn missing in index.html"
        assert 'Экспорт в Excel' in html, "'Экспорт в Excel' text missing in HTML"
        print("HTML contains corrExportExcelBtn button: OK")
        
    # JS
    with urllib.request.urlopen("http://127.0.0.1:8000/static/app.js") as resp:
        js = resp.read().decode("utf-8")
        assert 'function getCorrGradientBg' in js, "getCorrGradientBg missing in app.js"
        assert 'function getBetaGradientBg' in js, "getBetaGradientBg missing in app.js"
        assert 'function exportCorrMatrixToExcel' in js, "exportCorrMatrixToExcel missing in app.js"
        assert 'corrExportExcelBtn.addEventListener("click", exportCorrMatrixToExcel)' in js, "export click handler missing"
        assert 'getCorrGradientBg(item.correlation)' in js, "getCorrGradientBg call missing in renderCorrTableRows"
        assert 'getBetaGradientBg(item.beta)' in js, "getBetaGradientBg call missing in renderCorrTableRows"
        print("JS contains gradient helpers and export handler: OK")

    # CSS
    with urllib.request.urlopen("http://127.0.0.1:8000/static/style.css") as resp:
        css = resp.read().decode("utf-8")
        assert '.corr-export-btn' in css, ".corr-export-btn missing in style.css"
        assert '.corr-val-cell' in css, ".corr-val-cell missing in style.css"
        assert '.beta-val-cell' in css, ".beta-val-cell missing in style.css"
        print("CSS contains cell classes: OK")


if __name__ == "__main__":
    test_excel_export()
    test_html_and_js_served()
    print("ALL TESTS PASSED SUCCESSFULLY!")
